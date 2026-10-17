from app.domain.errors import ApplicationError
import base64
import csv
import io
from datetime import date
from uuid import uuid4

import pytest
from fastapi import HTTPException
from app.application.use_cases.employee_import import parse_file, build_preview, HEADERS
from app.domain.models import UserContext
from test_academy import api, Query, WORKER, MANAGER, OTHER, HR


def csv_content(rows):
    buffer = io.StringIO()
    csv.writer(buffer).writerows(rows)
    return base64.b64encode(buffer.getvalue().encode("utf-8-sig")).decode()


def row(**changes):
    return dict(
        employee_code="00001",
        full_name="Test Employee",
        email="test@example.test",
        company="PostTech",
        job_title="Worker",
        department="Warehouse",
        hired_at="2026-10-17",
        role="worker",
        manager_email="",
        **changes,
    )


def test_csv_bom_and_xlsx_dates_preserve_employee_codes():
    from openpyxl import Workbook

    value = row()
    parsed = parse_file("employees.csv", csv_content([HEADERS, list(value.values())]))
    assert parsed[0]["data"]["employee_code"] == "00001"
    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    ws.append([date(2026, 10, 17) if h == "hired_at" else value[h] for h in HEADERS])
    data = io.BytesIO()
    wb.save(data)
    parsed = parse_file("employees.xlsx", base64.b64encode(data.getvalue()).decode())
    assert parsed[0]["data"]["hired_at"] == "2026-10-17"
    assert parsed[0]["data"]["employee_code"] == "00001"


def test_import_rejects_formulas_unknown_fields_and_duplicate_headers():
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.append(HEADERS)
    values = list(row().values())
    values[1] = '=HYPERLINK("https://example.test")'
    ws.append(values)
    data = io.BytesIO()
    wb.save(data)
    with pytest.raises((HTTPException, ApplicationError)):
        parse_file("test.xlsx", base64.b64encode(data.getvalue()).decode())
    for headers in (HEADERS + ["role"], HEADERS + ["vacation_days_remaining"]):
        with pytest.raises((HTTPException, ApplicationError)):
            parse_file("test.csv", csv_content([headers, list(row().values())]))
    with pytest.raises((HTTPException, ApplicationError)):
        parse_file("test.csv", csv_content([HEADERS] + [list(row().values())] * 501))


def test_excel_rejects_far_away_cells_and_numeric_employee_codes():
    from openpyxl import Workbook

    for cell, value in [("XFD1", "Unexpected"), ("A2", 123)]:
        wb = Workbook()
        ws = wb.active
        ws.append(HEADERS)
        ws.append(list(row().values()))
        ws[cell] = value
        data = io.BytesIO()
        wb.save(data)
        with pytest.raises((HTTPException, ApplicationError)):
            parse_file("test.xlsx", base64.b64encode(data.getvalue()).decode())


def test_preview_shows_changes_and_prevents_access_change_and_duplicates():
    existing = {
        **row(),
        "id": WORKER,
        "manager_id": None,
        "updated_at": "2026-10-17T00:00:00Z",
    }
    incoming = {**row(), "department": "Delivery"}
    rows, errors = build_preview([{"line": 2, "data": incoming}], [existing], [])
    assert errors == [] and rows[0]["changes"] == [
        {"field": "department", "before": "Warehouse", "after": "Delivery"}
    ]
    for changed in (
        {"role": "hr_admin"},
        {"email": "other@example.test"},
        {"manager_email": "missing@example.test"},
    ):
        _, errors = build_preview(
            [{"line": 2, "data": {**incoming, **changed}}], [existing], []
        )
        assert errors
    rows, errors = build_preview(
        [{"line": 2, "data": row()}, {"line": 3, "data": row()}], [], []
    )
    assert len(errors) == 1 and len(rows) == 1
    _, errors = build_preview(
        [{"line": 2, "data": {**row(), "role": "hr_admin"}}], [], []
    )
    assert errors


@pytest.mark.parametrize("role", ["worker", "manager"])
@pytest.mark.parametrize(
    "method,path,payload",
    [
        ("GET", "/fulfillment", None),
        ("GET", "/tracks", None),
        ("GET", "/workflow-policy", None),
        ("POST", "/workflow-policy/run", None),
        ("GET", "/employee-import/template", None),
        (
            "POST",
            "/employee-import/preview",
            {"filename": "test.csv", "content": "YQ=="},
        ),
        ("POST", f"/employee-import/{uuid4()}/apply", None),
        ("POST", f"/tracks/{uuid4()}/enroll", {"employee_id": WORKER}),
        ("DELETE", f"/tracks/{uuid4()}", None),
        (
            "POST",
            f"/fulfillment/{uuid4()}",
            {"action": "complete", "reference": "Test"},
        ),
    ],
)
def test_admin_workflows_reject_non_hr(api, role, method, path, payload):
    client, db, actor = api
    actor["user"] = UserContext(
        id=WORKER if role == "worker" else MANAGER, auth_user_id=str(uuid4()), role=role
    )
    assert (
        client.request(method, "/api/v1/workspace" + path, json=payload).status_code
        == 403
    )
    assert not db.rpc_calls


def test_tracks_deny_foreign_employee_before_any_rpc(api):
    client, db, actor = api
    assert (
        client.get(f"/api/v1/workspace/my-tracks?employee_id={OTHER}").status_code
        == 403
    )
    actor["user"] = UserContext(id=MANAGER, auth_user_id=str(uuid4()), role="manager")
    assert (
        client.get(f"/api/v1/workspace/my-tracks?employee_id={OTHER}").status_code
        == 403
    )
    assert not db.rpc_calls


def test_policy_invalid_order_and_actor_injection_rejected(api):
    client, db, actor = api
    actor["user"] = UserContext(id=HR, auth_user_id=str(uuid4()), role="hr_admin")
    assert (
        client.put(
            "/api/v1/workspace/workflow-policy",
            json=dict(
                enabled=True, remind_hours=48, escalate_hours=24, course_notice_days=3
            ),
        ).status_code
        == 422
    )
    assert (
        client.post(
            f"/api/v1/workspace/tracks/{uuid4()}/enroll",
            json=dict(employee_id=WORKER, actor_id=HR),
        ).status_code
        == 422
    )


def test_summary_uses_authenticated_scope_without_calling_llm(api, monkeypatch):
    from app.application.use_cases import workflow_summary as workflow_summary
    from app.presentation.http.routes import ai as ai

    client, db, actor = api

    def neq(self, k, v):
        self.filters.append(lambda r: r.get(k) != v)
        return self

    def lt(self, k, v):
        self.filters.append(lambda r: r.get(k) is not None and r[k] < v)
        return self

    monkeypatch.setattr(Query, "neq", neq, raising=False)
    monkeypatch.setattr(Query, "lt", lt, raising=False)
    seen = []

    def scoped(user, **kwargs):
        seen.append(user.id)
        return []

    monkeypatch.setattr(workflow_summary, "list_requests", scoped)
    monkeypatch.setattr(
        ai,
        "answer_question",
        lambda *args, **kwargs: pytest.fail(
            "Operational data must not be sent to the model"
        ),
    )
    for a in db.tables["course_assignments"]:
        a["due_on"] = "2020-01-01"
    response = client.post("/api/v1/ai/ask", json={"question": "Что требует внимания?"})
    assert response.status_code == 200
    report = response.json()["report"]
    assert len(report["sections"]) == 2
    assert report["sections"][1]["count"] == 1
    assert "Other" not in response.text and seen == [WORKER]
