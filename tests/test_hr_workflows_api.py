import base64
import csv
import io
from datetime import date
from uuid import uuid4

import pytest
from fastapi import HTTPException
from app.services.employee_import import parse_file, build_preview, HEADERS
from app.schemas.common import UserContext
from test_academy import api, Query, WORKER, MANAGER, OTHER, HR


def csv_content(rows):
    buffer=io.StringIO();csv.writer(buffer).writerows(rows)
    return base64.b64encode(buffer.getvalue().encode('utf-8-sig')).decode()


def row(**changes):
    return {**dict(employee_code='00001',full_name='Test Employee',email='test@example.test',company='PostTech',job_title='Worker',department='Warehouse',hired_at='2026-10-17',role='worker',manager_email=''),**changes}


def test_csv_bom_and_xlsx_dates_preserve_employee_codes():
    from openpyxl import Workbook
    value=row()
    parsed=parse_file('employees.csv',csv_content([HEADERS,list(value.values())]))
    assert parsed[0]['data']['employee_code']=='00001'
    wb=Workbook();ws=wb.active;ws.append(HEADERS);ws.append([date(2026,10,17) if h=='hired_at' else value[h] for h in HEADERS]);data=io.BytesIO();wb.save(data)
    parsed=parse_file('employees.xlsx',base64.b64encode(data.getvalue()).decode())
    assert parsed[0]['data']['hired_at']=='2026-10-17'
    assert parsed[0]['data']['employee_code']=='00001'


def test_import_rejects_formulas_unknown_fields_and_duplicate_headers():
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.append(HEADERS);values=list(row().values());values[1]='=HYPERLINK("https://example.test")';ws.append(values);data=io.BytesIO();wb.save(data)
    with pytest.raises(HTTPException):parse_file('test.xlsx',base64.b64encode(data.getvalue()).decode())
    for headers in (HEADERS+['role'],HEADERS+['vacation_days_remaining']):
        with pytest.raises(HTTPException):parse_file('test.csv',csv_content([headers,list(row().values())]))
    with pytest.raises(HTTPException):parse_file('test.csv',csv_content([HEADERS]+[list(row().values())]*501))


def test_excel_rejects_far_away_cells_and_numeric_employee_codes():
    from openpyxl import Workbook
    for cell,value in [('XFD1','Unexpected'),('A2',123)]:
        wb=Workbook();ws=wb.active;ws.append(HEADERS);ws.append(list(row().values()));ws[cell]=value
        data=io.BytesIO();wb.save(data)
        with pytest.raises(HTTPException):parse_file('test.xlsx',base64.b64encode(data.getvalue()).decode())


def test_preview_shows_changes_and_prevents_access_change_and_duplicates():
    existing={**row(),'id':WORKER,'manager_id':None,'updated_at':'2026-10-17T00:00:00Z'}
    incoming={**row(),'department':'Delivery'}
    rows,errors=build_preview([{'line':2,'data':incoming}],[existing],[])
    assert errors==[] and rows[0]['changes']==[{'field':'department','before':'Warehouse','after':'Delivery'}]
    for changed in ({'role':'hr_admin'},{'email':'other@example.test'},{'manager_email':'missing@example.test'}):
        _,errors=build_preview([{'line':2,'data':{**incoming,**changed}}],[existing],[])
        assert errors
    rows,errors=build_preview([{'line':2,'data':row()},{'line':3,'data':row()}],[],[])
    assert len(errors)==1 and len(rows)==1
    _,errors=build_preview([{'line':2,'data':{**row(),'role':'hr_admin'}}],[],[])
    assert errors


def test_import_recognizes_existing_directors_but_does_not_create_privileged_access():
    director={**row(role='director'),'id':MANAGER,'manager_id':None,'updated_at':'2026-10-17T00:00:00Z'}
    rows,errors=build_preview([{'line':2,'data':row(role='director',department='Updated')}],[director],[])
    assert errors==[] and rows[0]['action']=='update'
    subordinate=row(employee_code='00002',email='worker@example.test',manager_email=director['email'])
    rows,errors=build_preview([{'line':2,'data':subordinate}],[director],[])
    assert errors==[] and rows[0]['manager_id']==MANAGER
    for role in ('director','hr_admin'):
        _,errors=build_preview([{'line':2,'data':row(role=role)}],[],[])
        assert errors


@pytest.mark.parametrize('role',['worker','manager','director'])
@pytest.mark.parametrize('method,path,payload',[
    ('GET','/fulfillment',None),('GET','/tracks',None),('GET','/workflow-policy',None),
    ('POST','/workflow-policy/run',None),('GET','/employee-import/template',None),
    ('POST','/employee-import/preview',{'filename':'test.csv','content':'YQ=='}),
    ('POST',f'/employee-import/{uuid4()}/apply',None),
    ('POST',f'/tracks/{uuid4()}/enroll',{'employee_id':WORKER}),
    ('DELETE',f'/tracks/{uuid4()}',None),
    ('POST',f'/fulfillment/{uuid4()}',{'action':'complete','reference':'Test'}),
])
def test_admin_workflows_reject_non_hr(api,role,method,path,payload):
    client,db,actor=api;actor['user']=UserContext(id=WORKER if role=='worker' else MANAGER,auth_user_id=str(uuid4()),role=role)
    assert client.request(method,'/api/v1/workspace'+path,json=payload).status_code==403
    assert not db.rpc_calls


def test_tracks_deny_foreign_employee_before_any_rpc(api):
    client,db,actor=api
    assert client.get(f'/api/v1/workspace/my-tracks?employee_id={OTHER}').status_code==403
    actor['user']=UserContext(id=MANAGER,auth_user_id=str(uuid4()),role='manager')
    assert client.get(f'/api/v1/workspace/my-tracks?employee_id={OTHER}').status_code==403
    assert not db.rpc_calls


def test_policy_invalid_order_and_actor_injection_rejected(api):
    client,db,actor=api;actor['user']=UserContext(id=HR,auth_user_id=str(uuid4()),role='hr_admin')
    assert client.put('/api/v1/workspace/workflow-policy',json=dict(enabled=True,remind_hours=48,escalate_hours=24,course_notice_days=3)).status_code==422
    assert client.post(f'/api/v1/workspace/tracks/{uuid4()}/enroll',json=dict(employee_id=WORKER,actor_id=HR)).status_code==422


def test_summary_uses_authenticated_scope_without_calling_llm(api,monkeypatch):
    from app.services import workflow_summary
    from app.api.routes import ai
    client,db,actor=api
    def neq(self,k,v):self.filters.append(lambda r:r.get(k)!=v);return self
    def lt(self,k,v):self.filters.append(lambda r:r.get(k) is not None and r[k]<v);return self
    monkeypatch.setattr(Query,'neq',neq,raising=False);monkeypatch.setattr(Query,'lt',lt,raising=False)
    seen=[]
    def scoped(user,**kwargs):seen.append(user.id);return []
    monkeypatch.setattr(workflow_summary,'list_requests',scoped)
    monkeypatch.setattr(ai,'answer_question',lambda *args,**kwargs:pytest.fail('Operational data must not be sent to the model'))
    for a in db.tables['course_assignments']:a['due_on']='2020-01-01'
    response=client.get('/api/v1/workspace/summary')
    assert response.status_code==200
    report=response.json();assert len(report['sections'])==2
    assert report['sections'][1]['count']==1
    assert 'Other' not in response.text and seen==[WORKER]


@pytest.mark.parametrize('role,expected_courses',[('worker',1),('manager',1),('director',1),('hr_admin',2)])
def test_operational_summary_preserves_all_four_role_scopes(api,monkeypatch,role,expected_courses):
    from app.services import workflow_summary
    client,db,actor=api
    actor_id=WORKER if role=='worker' else HR if role=='hr_admin' else MANAGER
    actor['user']=UserContext(id=actor_id,auth_user_id=str(uuid4()),role=role)
    def neq(self,k,v):self.filters.append(lambda r:r.get(k)!=v);return self
    def lt(self,k,v):self.filters.append(lambda r:r.get(k) is not None and r[k]<v);return self
    def either(self,expression):
        terms=[term.split('.',2) for term in expression.split(',')]
        assert all(op=='eq' for _,op,_ in terms)
        self.filters.append(lambda row:any(row.get(key)==value for key,_,value in terms));return self
    monkeypatch.setattr(Query,'neq',neq,raising=False)
    monkeypatch.setattr(Query,'lt',lt,raising=False)
    monkeypatch.setattr(Query,'or_',either,raising=False)
    db.tables['request_fulfillment']=[]
    for assignment in db.tables['course_assignments']:assignment['due_on']='2020-01-01'
    def scoped(user,**kwargs):
        return [dict(id=str(uuid4()),employee_id=WORKER,status='pending_hr',request_type='advance',
                     created_at='2020-01-01',can_review=user.role=='hr_admin')]
    monkeypatch.setattr(workflow_summary,'list_requests',scoped)
    response=client.get('/api/v1/workspace/summary')
    assert response.status_code==200,response.text
    sections=response.json()['sections']
    assert sections[0]['count']==(1 if role in ('worker','hr_admin') else 0)
    assert sections[1]['count']==expected_courses
    assert ('Other' in response.text)==(role=='hr_admin')
